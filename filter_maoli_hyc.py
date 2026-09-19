#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
filter_maoli_hyc.py —— 把「全公司」门店毛利明细表过滤成「华为华阳城合作店」专用明细。

为什么必须有这一步：
  李家村用店长账号(18161914293)导出，返回天然只有李家村，所以 calc_data.py /
  write_xs_xml.py 都**不做门店过滤**（这是历史设计，不是 bug）。
  华阳城必须用经理账号(18591910491)，该账号导出的是**全公司 17 店**明细
  （实测 10815 行）。若不过滤直接喂管线：
    · write_xs_xml.py 会把 17 家店全部写进底表 XS → 底表 SUMIFS 全部算错；
    · calc_data.py 会把 17 家店当成一家店复算 → 看板数字全部虚高。
  故在管线最前面加这一道过滤，让 calc_data / write_xs_xml **保持与 shop 逐字一致**。

过滤口径（实测 2026-09-19，两列完全一致、零分歧）：
  保留「库区」或「销售出库单门店」包含 "华阳城" 的行。

用法：
  python filter_maoli_hyc.py <源明细.xlsx> [输出.xlsx]
"""
import os
import sys
import datetime
import warnings

from openpyxl import load_workbook, Workbook

STORE_KEY = "华阳城"
HEADERS = ["出库单号", "单据类型", "出库日期", "商品分类", "商品sku分类", "商品SKU编码",
           "商品名称", "入库属性", "数量", "单价", "原价", "折扣价", "金额", "毛利",
           "SO激励", "业务员", "库区", "销售出库单门店", "销售成本"]
DEF_OUT_DIR = "/Users/mac/.local/share/TeleAgent/playwright-mcp/hyc"
DEF_OUT = os.path.join(DEF_OUT_DIR, "华阳城门店毛利明细表-华为终端.xlsx")


def main():
    if len(sys.argv) < 2:
        sys.exit("用法: filter_maoli_hyc.py <源明细.xlsx> [输出.xlsx]")
    src = sys.argv[1]
    out = sys.argv[2] if len(sys.argv) > 2 else DEF_OUT
    if not os.path.exists(src):
        sys.exit(f"❌ 找不到源明细: {src}")

    warnings.filterwarnings("ignore", message="Workbook contains no default style")
    wb = load_workbook(src, read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    hdr = [str(h).strip() if h is not None else "" for h in next(it)]
    if hdr[:19] != HEADERS:
        sys.exit(f"❌ 源明细表头不符（期望 19 列）\n  实际: {hdr[:19]}")
    i_kq, i_md = 16, 17   # 库区 / 销售出库单门店

    keep, total = [], 0
    for r in it:
        if not r or not r[0] or not str(r[0]).strip():
            continue
        total += 1
        kq = STORE_KEY in str(r[i_kq] or "")
        md = STORE_KEY in str(r[i_md] or "")
        if kq or md:
            keep.append(r[:19])

    if not keep:
        sys.exit(f"❌ 过滤后 0 行（源 {total} 行中无「{STORE_KEY}」）—— 检查门店名")

    os.makedirs(os.path.dirname(out), exist_ok=True)
    nwb = Workbook(write_only=True)
    nws = nwb.create_sheet("门店毛利明细表-华为终端")
    nws.append(HEADERS)
    for r in keep:
        nws.append(["" if v is None else v for v in r])
    nwb.save(out)

    # 校验：行数 + 毛利合计 + 人员分布
    import collections
    emp = collections.Counter(str(r[15] or "") for r in keep)
    gross = 0.0
    for r in keep:
        try:
            gross += float(str(r[13]).replace(",", ""))
        except (TypeError, ValueError):
            pass
    print(f"→ 源 {total} 行 → 保留「{STORE_KEY}」{len(keep)} 行")
    print(f"  毛利合计 ¥{gross:,.2f}")
    print(f"  业务员分布: {dict(emp)}")
    print(f"✓ 已写出: {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
