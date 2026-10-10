#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
hyc_oct_surgery.py —— 华阳城底表 9月→10月 手术（2026-10-11 18号）
把《华阳城9月任务进度.xlsx》克隆为《华阳城10月任务进度.xlsx》，
核心变化：任务人 4人 → 5人（新增 乔玉宇），填入10月任务数字。

布局变化（华阳城销售 sheet）：
  P1: 表头2-3不变 | 人 4-8(5人) | 田蕊9 | 合计10
  P2: 表头12-14   | 人 15-19    | 田蕊20 | 合计21
  P3: 标题28 表头29-30 | 人 31-35 | 田蕊36 | 合计37
  P4: 标题39 表头40-41 | 人 42-46 | 合计47
（月度任务 sheet：人4-8，合计9）

用法： python3 hyc_oct_surgery.py
输入： /Users/mac/Desktop/华阳城销售/华阳城9月任务进度.xlsx（只读，绝不修改）
输出： /Users/mac/Desktop/华阳城销售/华阳城10月任务进度.xlsx
"""
import re, sys, json
from copy import copy
import openpyxl
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.utils import column_index_from_string, get_column_letter

SRC = "/Users/mac/Desktop/华阳城销售/华阳城9月任务进度.xlsx"
DST = "/Users/mac/Desktop/华阳城销售/华阳城10月任务进度.xlsx"

PEOPLE = ["张梅B", "李俊琪", "王莹莹B", "张赫桐", "乔玉宇"]   # 月任务行 4-8
# 10月任务：C收入 D毛利 E手机 F增值服务 G合约 H(PC/电脑) I平板 J音频 K穿戴 L(HD/智慧屏)
OCT = {
    "张梅B":   dict(C=600000, D=78000, E=60, F=18000, G=3, H=4, I=5, J=9, K=4, L=0),
    "李俊琪":  dict(C=600000, D=78000, E=60, F=18000, G=3, H=4, I=5, J=9, K=4, L=1),
    "王莹莹B": dict(C=600000, D=78000, E=60, F=18000, G=3, H=4, I=5, J=9, K=4, L=1),
    "张赫桐":  dict(C=600000, D=78000, E=60, F=18000, G=3, H=4, I=5, J=9, K=4, L=0),
    "乔玉宇":  dict(C=400000, D=58000, E=60, F=11000, G=1, H=2, I=5, J=9, K=0, L=0),
}
QD_TASK = 41400          # 渠道挂账任务/人（沿用9月标准，待晨哥确认）
KEEP_COLS = {"M": 0, "N": 4, "O": 4}   # 摄影课程/考核机/滞销：10月表未给，沿用9月口径

CELLREF = re.compile(r'(?<![A-Z0-9_."])([A-Z]{1,3})(\d{1,3})(?![0-9(])')
# 带表名前缀的引用（个人sheet/店长/薪资用：只改跨表引用，不动本表本地引用）
PREFREF = re.compile(r'((?:华阳城销售|月度任务)![A-Z]{1,2})(\d{1,3})')

def shift_refs(f, mapping):
    """按 {旧行:新行} 映射改写公式里的单元格引用行号（含范围两端点、跨表引用）。"""
    def rep(m):
        col, row = m.group(1), int(m.group(2))
        return col + str(mapping.get(row, row))
    return CELLREF.sub(rep, f)

def shift_pref(f, mapping):
    """只改写带 华阳城销售!/月度任务! 前缀的引用行号。"""
    def rep(m):
        return m.group(1) + str(mapping.get(int(m.group(2)), int(m.group(2))))
    return PREFREF.sub(rep, f)

def shift_afref(ref, mapping):
    """数组公式 ref（如 S14 或 S14:S14）改写。"""
    return shift_refs(ref, mapping)

def main():
    wb = openpyxl.load_workbook(SRC, data_only=False)
    ws = wb["华阳城销售"]

    # ---------- 0. 捕获模板（值/公式/样式/行高） ----------
    def snap(row, cmax=41):
        d = {}
        for c in range(1, cmax + 1):
            cell = ws.cell(row, c)
            v = cell.value
            if isinstance(v, ArrayFormula):
                d[c] = {"af": True, "ref": v.ref, "text": v.text, "style": copy(cell._style)}
            elif v is not None:
                d[c] = {"v": v, "style": copy(cell._style)}
            else:
                d[c] = {"style": copy(cell._style)}
        return d

    T = {r: snap(r) for r in [4, 8, 9, 11, 12, 13, 14, 18, 19,
                              25, 26, 27, 28, 32, 33, 35, 36, 37, 38, 39, 40, 41, 42]}
    BLANK = snap(20)                    # 旧空行样式（行20-24是空隔行）
    H = {r: ws.row_dimensions[r].height for r in list(T) + [34]}

    # ---------- 1. 清空 4-47 行（值），解除该区域全部合并 ----------
    old_merges = [str(m) for m in ws.merged_cells.ranges]
    for mref in old_merges:
        if int(re.search(r"\d+", mref).group()) >= 4:
            ws.unmerge_cells(mref)
    for r in range(4, 48):
        for c in range(1, 42):
            ws.cell(r, c).value = None

    def put(row, col, spec, mapping=None, name=None):
        cell = ws.cell(row, col)
        cell._style = copy(spec["style"])
        if "af" in spec:
            text = spec["text"]
            ref = spec["ref"]
            if name:
                text = text.replace('"张梅B"', '"%s"' % name)
            if mapping:
                text = shift_refs(text, mapping)
                ref = shift_afref(ref, mapping)
            cell.value = ArrayFormula(ref=ref, text=text)
        elif "v" in spec:
            v = spec["v"]
            if isinstance(v, str) and (v.startswith("=") or "!" in v):
                if name:
                    v = v.replace('"张梅B"', '"%s"' % name)
                cell.value = shift_refs(v, mapping) if mapping else v
            else:
                cell.value = v
        return cell

    def rowheight(row, h):
        if h:
            ws.row_dimensions[row].height = h

    # ---------- 2. P1 块（人4-8 田蕊9 合计10） ----------
    for i, nm in enumerate(PEOPLE):
        r = 4 + i
        mp = {4: 4 + i}
        for c, spec in T[4].items():
            put(r, c, spec, mp, nm)
        rowheight(r, H[4])
    for c, spec in T[8].items():
        put(9, c, spec, {8: 9}, "田蕊")
    rowheight(9, H[8])
    for c, spec in T[9].items():
        put(10, c, spec, {9: 10, 8: 9, 4: 4})
    rowheight(10, H[9])

    # ---------- 3. P2 块（表头12-14 人15-19 田蕊20 合计21） ----------
    for new_r, old_r in [(12, 11), (13, 12), (14, 13)]:
        for c, spec in T[old_r].items():
            put(new_r, c, spec)
        rowheight(new_r, H[old_r])
    for i, nm in enumerate(PEOPLE):
        r = 15 + i
        mp = {14: 15 + i, 4: 4 + i}
        for c, spec in T[14].items():
            put(r, c, spec, mp, nm)
        rowheight(r, H[14])
    for c, spec in T[18].items():
        put(20, c, spec, {18: 20, 8: 9}, "田蕊")
    rowheight(20, H[18])
    for c, spec in T[19].items():
        put(21, c, spec, {19: 21, 18: 20, 14: 15, 9: 10})
    rowheight(21, H[19])

    # ---------- 4. P3 块（标题28 表头29-30 人31-35 田蕊36 合计37） ----------
    for new_r, old_r in [(28, 25), (29, 26), (30, 27)]:
        for c, spec in T[old_r].items():
            put(new_r, c, spec)
        rowheight(new_r, H[old_r])
    for i, nm in enumerate(PEOPLE):
        r = 31 + i
        mp = {28: 31 + i, 4: 4 + i}
        for c, spec in T[28].items():
            put(r, c, spec, mp, nm)
        rowheight(r, H[28])
    for c, spec in T[32].items():
        put(36, c, spec, {32: 36, 8: 9}, "田蕊")
    rowheight(36, H[32])
    for c, spec in T[33].items():
        put(37, c, spec, {33: 37, 28: 31, 32: 36})
    rowheight(37, H[33])

    # ---------- 5. P4 块（标题39 表头40-41 人42-46 合计47） ----------
    for new_r, old_r in [(39, 35), (40, 36), (41, 37)]:
        for c, spec in T[old_r].items():
            put(new_r, c, spec)
        rowheight(new_r, H[old_r])
    for i, nm in enumerate(PEOPLE):
        r = 42 + i
        if i < 4:
            src = 38 + i          # 用各自旧模板行（引用行=4+i，恒等映射）
            mp = {4: 4 + i}
        else:                      # 乔玉宇：用 r38 模板，P1引用 4→8
            src = 38
            mp = {4: 8}
        for c, spec in T[src].items():
            put(r, c, spec, mp, nm)
        rowheight(r, H[38])
    for c, spec in T[42].items():
        put(47, c, spec, {42: 47, 38: 42, 41: 46})
    rowheight(47, H[42])

    # 空隔行行高与样式（新11 / 22-27 / 38）
    for br in [11] + list(range(22, 28)) + [38]:
        for c, spec in BLANK.items():
            ws.cell(br, c)._style = copy(spec["style"])
    rowheight(38, H.get(34))

    # ---------- 6. 重建合并单元格 ----------
    def map_row(r):
        if r <= 3:
            return r
        if 4 <= r <= 7:
            return r
        if r == 8:
            return 9          # 旧P1田蕊
        if r == 9:
            return 10         # 旧P1合计
        if 11 <= r <= 17:
            return r + 1      # P2 表头+人
        if r == 18:
            return 20         # 旧P2田蕊
        if r == 19:
            return 21         # 旧P2合计
        if 25 <= r <= 32:
            return r + 3      # P3
        if r == 33:
            return 37         # 旧P3合计
        if 35 <= r <= 41:
            return r + 4      # P4
        if r == 42:
            return 47         # 旧P4合计
        return None

    for mref in old_merges:
        a = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", mref)
        if not a:
            continue
        c1, r1, c2, r2 = a.group(1), int(a.group(2)), a.group(3), int(a.group(4))
        if r1 < 4:
            continue                      # 表头区未动，本来就是合并状态
        nr1, nr2 = map_row(r1), map_row(r2)
        if nr1 and nr2:
            ws.merge_cells("%s%d:%s%d" % (c1, nr1, c2, nr2))
    # P2 各人行补 AJ:AK 合并（旧结构每人行都有）
    for i in range(5):
        ws.merge_cells("AJ%d:AK%d" % (15 + i, 15 + i))

    # ---------- 7. 月度任务 sheet ----------
    tk = wb["月度任务"]
    tk["A2"] = "华为华阳城2026年10月"
    style_person = [copy(tk.cell(4, c)._style) for c in range(1, 16)]
    style_total = [copy(tk.cell(8, c)._style) for c in range(1, 16)]
    for i, nm in enumerate(PEOPLE):
        r = 4 + i
        vals = {"B": nm}
        vals.update({k: v for k, v in OCT[nm].items()})
        vals.update(KEEP_COLS)          # M摄影课程 N考核机 O滞销 沿用
        for c in range(1, 16):
            cell = tk.cell(r, c)
            cell._style = copy(style_person[c - 1])
            L = get_column_letter(c)
            cell.value = vals.get(L)
    totalvals = {"A": "合计", "O": "合计"}
    for c in range(1, 16):
        cell = tk.cell(9, c)
        cell._style = copy(style_total[c - 1])
        L = get_column_letter(c)
        if L in totalvals:
            cell.value = totalvals[L]
        elif "C" <= L <= "N":
            cell.value = "=SUM(%s4:%s8)" % (L, L)
        else:
            cell.value = None

    # ---------- 8. 渠道挂账 sheet：插乔玉宇行 ----------
    qd = wb["渠道挂账"]
    qd.insert_rows(8)
    # 新行8 = 乔玉宇（克隆样式 from 行7）
    c4_full = qd.cell(4, 3).value     # 完整 SUMIFS 公式（引用 A4）
    for c in range(1, 6):
        qd.cell(8, c)._style = copy(qd.cell(7, c)._style)
    qd.cell(8, 1).value = "乔玉宇"
    qd.cell(8, 2).value = QD_TASK
    qd.cell(8, 3).value = shift_refs(c4_full, {4: 8})
    qd.cell(8, 4).value = "=C8-B8"
    qd.cell(8, 5).value = "=C8/B8"
    # 旧田蕊行(现9)：A8→A9
    for c in range(1, 6):
        v = qd.cell(9, c).value
        if isinstance(v, str) and v.startswith("="):
            qd.cell(9, c).value = shift_refs(v, {8: 9})
    # 旧合计行(现10)：9→10, SUM端点 8→9
    for c in range(1, 6):
        v = qd.cell(10, c).value
        if isinstance(v, str) and v.startswith("="):
            qd.cell(10, c).value = shift_refs(v, {9: 10, 8: 9})

    # ---------- 9. 个人sheet / 店长 / 薪资明细 引用行改写（只动跨表引用） ----------
    # 乔玉宇 sheet：先克隆原始张梅B（未改写前），再改自己的跨表引用
    src = wb["张梅B"]
    joe = wb.copy_worksheet(src)
    joe.title = "乔玉宇"
    joe["A1"] = "乔玉宇-绩效考评表"
    mp_joe = {4: 8, 14: 19}      # 月任务 4→8；华阳城销售 P1 4→8、P2 14→19
    for row in joe.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.value = shift_pref(cell.value, mp_joe)
    # 张梅B(i0) 李俊琪(i1) 王莹莹B(i2) 张赫桐(i3)：华阳城销售 P2 行引用 14-17 → +1
    for i, nm in enumerate(PEOPLE[:4]):
        w = wb[nm]
        mp = {k: k + 1 for k in range(14, 18)}
        for row in w.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    cell.value = shift_pref(cell.value, mp)
    # 店长：月度任务 8→9；华阳城销售 9→10, 19→21
    dz = wb["店长"]
    for row in dz.iter_rows():
        for cell in row:
            v = cell.value
            if not (isinstance(v, str) and v.startswith("=")):
                continue
            if "月度任务!" in v:
                v = re.sub(r'(月度任务!([A-Z]{1,2}))8\b', r'\g<1>9', v)
            if "华阳城销售!" in v:
                v = shift_pref(v, {9: 10, 19: 21})
            cell.value = v
    # 薪资明细：华阳城销售 14-16 → +1
    xz = wb["薪资明细"]
    for row in xz.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and v.startswith("=") and "华阳城销售!" in v:
                cell.value = shift_pref(v, {14: 15, 15: 16, 16: 17})

    # ---------- 10. 打开即全量重算 + 保存 ----------
    wb.calculation.fullCalcOnLoad = True
    wb.save(DST)
    print("✅ 生成:", DST)

if __name__ == "__main__":
    main()
