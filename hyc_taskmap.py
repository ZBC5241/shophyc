#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
hyc_taskmap.py —— 华阳城底表「结构自动发现」共享模块（2026-10-11 动态化改造）

解决什么问题：
    每月任务更新时，底表结构可能变（人员增减/行偏移/块移动），
    过去 calc_data / build_data / verify_hyc / perf_score / write_xs_xml / merge_qudao
    各自硬编码行号（range(4,9)、+11、田蕊=20 ……），每月都要手工改 6 个脚本。
    现在全部改为从本模块动态发现，脚本零硬编码。

原理（不依赖公式缓存值，纯公式文本反向解析）：
    1) 「月度任务」sheet：B 列从 r4 向下读姓名，遇「合计」停 → task_people / task_rows
    2) 「华阳城销售」sheet：A 列逐行扫，人员行是公式「=月度任务!B{n}」，
       从公式文本提取 n 即得 {月任务行: 底表行} 映射；
       字面姓名（田蕊等不背任务者）→ extra；「合计」→ 块边界。
    3) 按合计行切成连续块 → P1(业绩)/P2(乐机收)/P3(每日完成)/P4(每日缺口)。

用法：
    import hyc_taskmap
    tm = hyc_taskmap.discover("/Users/mac/Desktop/华阳城销售/华阳城10月任务进度.xlsx")
    # tm["task_people"] / tm["task_rows"] / tm["people_order"]
    # tm["P1_ROWS"] tm["P1_TIANRUI"] tm["P1_TOTAL"] …… P2/P3/P4 同理
"""

import re
import openpyxl

MONTH_RE = re.compile(r"^\s*=?\s*'?月度任务'?!B(\d+)\s*$")

def _month_sheet(wb):
    names = [s for s in wb.sheetnames if s.endswith("月任务") or s.endswith("度任务")]
    if not names:
        raise RuntimeError("底表里找不到「月度任务」sheet: %s" % wb.sheetnames)
    return names[0]

def discover(xlsx):
    """解析底表结构，返回全部行映射。结构异常时抛异常（宁可失败不可猜错）。"""
    wbf = openpyxl.load_workbook(xlsx, data_only=False)
    tk = wbf[_month_sheet(wbf)]

    # ---- 1. 月度任务人员名单：B 列 r4 起，遇「合计」停 ----
    task_people, task_rows = [], {}
    for r in range(4, 41):
        v = tk.cell(r, 2).value
        s = str(v).strip() if v is not None else ""
        if s == "合计":
            break
        if s:
            task_people.append(s)
            task_rows[s] = r
    if not task_people:
        raise RuntimeError("月度任务 B 列没读到任何人员姓名")

    # ---- 2. 华阳城销售 A 列扫描，按「合计」分块 ----
    # A 列形态：P1 =「=月度任务!B{n}」；P2~P4 =「=A{P1行}」本表引用；特殊行 = 字面姓名
    ws = wbf["华阳城销售"]
    raw = {}
    for r in range(1, 81):
        v = ws.cell(r, 1).value
        if v is not None:
            raw[r] = str(v).strip()

    memo = {}                                   # row -> ("ref", 月任务行) | ("name", 姓名)
    def parse(r):
        if r in memo:
            return memo[r]
        memo[r] = None                          # 防循环引用
        s = raw.get(r, "")
        m = MONTH_RE.match(s)
        if m:
            memo[r] = ("ref", int(m.group(1)))
        elif s.startswith("=A"):
            m2 = re.match(r"^=?A(\d+)$", s)
            if m2:
                memo[r] = parse(int(m2.group(1)))
        elif s and not s.startswith("="):
            memo[r] = ("name", s)
        return memo[r]

    blocks, cur = [], None
    for r in sorted(raw):
        kind = parse(r)
        if kind and kind[0] == "ref":
            if cur is None:
                cur = {"refs": {}, "extra": {}, "total": None}
            cur["refs"][kind[1]] = r
        elif raw.get(r) == "合计" and cur is not None:
            cur["total"] = r
            blocks.append(cur)
            cur = None
        elif kind and kind[0] == "name" and cur is not None:
            cur["extra"][kind[1]] = r
    if cur is not None:
        raise RuntimeError("华阳城销售 A 列扫到最后一块没有「合计」行")
    if len(blocks) != 4:
        raise RuntimeError("期望 4 个块(P1~P4)，实际扫到 %d 个" % len(blocks))

    # ---- 3. 组装映射：块内 refs 的 key 就是月任务行号，直接对上姓名 ----
    inv = {v: k for k, v in task_rows.items()}          # 月任务行号 → 姓名
    out = {
        "task_people": task_people,
        "task_rows": dict(task_rows),
        "people_order": None, "helper": None,
        "P1_ROWS": {}, "P1_TIANRUI": None, "P1_TOTAL": None,
        "P2_ROWS": {}, "P2_TIANRUI": None, "P2_TOTAL": None,
        "P3_ROWS": {}, "P3_TIANRUI": None, "P3_TOTAL": None,
        "P4_ROWS": {}, "P4_TIANRUI": None, "P4_TOTAL": None,
    }
    helper_name = None
    for i, b in enumerate(blocks, 1):
        rows = {}
        for n, r in b["refs"].items():
            if n not in inv:
                raise RuntimeError("P%d 块引用月任务行 %d，但月度任务表无此人" % (i, n))
            rows[inv[n]] = r
        tianrui = None
        for name, r in b["extra"].items():
            if name not in rows:
                tianrui = (name, r)
        out["P%d_ROWS" % i] = rows
        out["P%d_TIANRUI" % i] = tianrui
        out["P%d_TOTAL" % i] = b["total"]
        if i == 1 and tianrui:
            helper_name = tianrui[0]
    out["helper"] = helper_name or "田蕊"
    out["people_order"] = task_people + [out["helper"]]
    # 田蕊（helper）不背月任务，但 P1 块可能有它的行
    if helper_name and helper_name not in out["P1_ROWS"] and out["P1_TIANRUI"]:
        out["P1_ROWS"][helper_name] = out["P1_TIANRUI"][1]
    return out

if __name__ == "__main__":
    import sys, json
    tm = discover(sys.argv[1] if len(sys.argv) > 1 else
                  "/Users/mac/Desktop/华阳城销售/华阳城10月任务进度.xlsx")
    print(json.dumps(tm, ensure_ascii=False, indent=1))
