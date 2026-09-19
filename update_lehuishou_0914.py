#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# 更新《华阳城9月任务进度.xlsx》「华阳城销售」sheet 乐机收手工项（T=单量 U=公司净利）
# 行: 14=张梅B 15=李俊琪 16=王莹莹B 17=张赫桐 18=田蕊
# 安全: 写前/写后校验 T19/U19 仍是 =SUM 公式; 不碰 V 列(增值)与太力回收.
import argparse, json, sys
import openpyxl

ROWS = {"张梅B": 14, "李俊琪": 15, "王莹莹B": 16, "张赫桐": 17, "田蕊": 18}


def num(v):
    if v is None:
        return 0.0
    if isinstance(v, str):
        v = v.replace(",", "").strip()
        if v == "":
            return 0.0
    try:
        return float(v)
    except Exception:
        return 0.0


def write(xlsx, data, dry_run=False, verbose=True):
    wb = openpyxl.load_workbook(xlsx, data_only=False)
    if "华阳城销售" not in wb.sheetnames:
        sys.exit("✗ 找不到 sheet「华阳城销售」")
    ws = wb["华阳城销售"]

    # 校验合计公式 (T19/U19 必须是 SUM; V19 当前为空, 不强制)
    for c, name in ((20, "T"), (21, "U")):
        f = ws.cell(19, c).value
        if not (isinstance(f, str) and f.upper().startswith("=SUM")):
            sys.exit(f"✗ {name}19 不是 SUM 公式（当前={f!r}），终止以防破坏合计")

    old = {n: (ws.cell(r, 20).value, ws.cell(r, 21).value) for n, r in ROWS.items()}
    for n, d in data.items():
        if n not in ROWS:
            print(f"  ⚠ 跳过未知人员 {n}")
            continue
        r = ROWS[n]
        if "orders" in d and d["orders"] is not None:
            ws.cell(r, 20).value = num(d["orders"])   # T 单量
        if "amount" in d and d["amount"] is not None:
            ws.cell(r, 21).value = num(d["amount"])    # U 公司净利
        if verbose:
            o = d.get("orders"); a = d.get("amount")
            print(f"  {n}: T {old[n][0]}→{o}   U {old[n][1]}→{a}")

    # 写后校验
    for c, name in ((20, "T"), (21, "U")):
        f = ws.cell(19, c).value
        if not (isinstance(f, str) and f.upper().startswith("=SUM")):
            sys.exit(f"✗ 写入后 {name}19 被破坏（{f!r}）")

    if dry_run:
        if verbose:
            print("（dry-run）未落盘")
        return

    wb.save(xlsx)
    if verbose:
        print("✓ 已写入:", xlsx)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--data", required=True, help='JSON')
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    data = json.loads(args.data)
    if not data:
        sys.exit("✗ 无数据")
    write(args.xlsx, data, dry_run=args.dry_run)


if __name__ == "__main__":
    main()
